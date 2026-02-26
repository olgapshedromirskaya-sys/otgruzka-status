import csv
import io
import logging
from datetime import datetime
from pathlib import Path
from typing import Optional
from urllib.parse import quote_plus

from apscheduler.schedulers.asyncio import AsyncIOScheduler
from fastapi import Depends, FastAPI, HTTPException, Query, Request
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from openpyxl import Workbook
from sqlalchemy import text
from sqlalchemy.orm import Session

from app.auth import build_signed_init_data, extract_telegram_id_from_init_data
from app.config import settings
from app.db import get_session, init_db, session_scope
from app.enums import Marketplace
from app.schemas import (
    DashboardSummary,
    OrdersResponse,
    SettingsRead,
    SettingsUpdate,
    StatusCatalogItem,
    SyncReport,
)
from app.services import (
    build_summary,
    ensure_owner_user,
    export_rows,
    get_settings,
    list_orders,
    marketplace_catalog,
    save_settings,
    status_catalog,
    sync_orders_from_marketplaces,
)

logger = logging.getLogger(__name__)
logging.basicConfig(level=logging.INFO)

app = FastAPI(title=settings.app_name, version="0.2.0")
static_dir = Path(__file__).resolve().parent / "static"

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)

app.mount("/static", StaticFiles(directory=static_dir), name="static")

bot_instance = None
dp_instance = None
scheduler: AsyncIOScheduler | None = None


def _extract_init_data(request: Request) -> str:
    return (
        request.headers.get("X-Telegram-Init-Data")
        or request.query_params.get("init_data")
        or request.query_params.get("tgWebAppData")
        or ""
    )


def _verify_webhook_init_data(request: Request) -> None:
    init_data = _extract_init_data(request)
    if not init_data:
        raise HTTPException(
            status_code=401,
            detail="Требуется авторизация Telegram initData",
        )

    try:
        extract_telegram_id_from_init_data(init_data, settings.bot_token)
    except ValueError as exc:
        raise HTTPException(
            status_code=401,
            detail=f"Ошибка авторизации webhook: {exc}",
        ) from exc


def _build_webhook_url() -> str:
    base_url = settings.webapp_url.rstrip("/")
    bot_id_raw = settings.bot_token.split(":", maxsplit=1)[0]
    try:
        bot_id = int(bot_id_raw)
    except ValueError as exc:
        raise RuntimeError("Некорректный BOT_TOKEN: не удалось определить bot_id") from exc

    signed_init_data = build_signed_init_data(settings.bot_token, bot_id)
    return f"{base_url}/webhook?init_data={quote_plus(signed_init_data)}"


async def _scheduled_sync_job() -> None:
    report = await sync_orders_from_marketplaces()
    logger.info(
        "Плановая синхронизация: WB=%s Ozon=%s обработано=%s",
        report.wb_received,
        report.ozon_received,
        report.processed_orders,
    )


def _start_scheduler() -> None:
    global scheduler
    if scheduler and scheduler.running:
        return

    scheduler = AsyncIOScheduler(timezone=settings.timezone)
    scheduler.add_job(
        _scheduled_sync_job,
        trigger="interval",
        minutes=15,
        id="wb-ozon-fbs-sync",
        max_instances=1,
        coalesce=True,
        misfire_grace_time=600,
    )
    scheduler.start()
    logger.info("APScheduler запущен (интервал: 15 минут)")


@app.on_event("startup")
async def startup_event() -> None:
    global bot_instance, dp_instance
    init_db()
    with session_scope() as session:
        ensure_owner_user(session)
    _start_scheduler()

    if settings.bot_token and settings.webapp_url:
        from app.bot import get_bot, get_dispatcher

        bot_instance = get_bot()
        dp_instance = get_dispatcher()
        webhook_url = _build_webhook_url()
        await bot_instance.delete_webhook(drop_pending_updates=True)
        await bot_instance.set_webhook(webhook_url)
        logger.info("Webhook set to %s/webhook", settings.webapp_url.rstrip("/"))


@app.on_event("shutdown")
async def shutdown_event() -> None:
    if scheduler and scheduler.running:
        scheduler.shutdown(wait=False)
    if bot_instance:
        await bot_instance.session.close()


@app.post("/webhook")
async def telegram_webhook(request: Request) -> dict:
    from aiogram.types import Update

    _verify_webhook_init_data(request)
    if bot_instance and dp_instance:
        data = await request.json()
        update = Update.model_validate(data)
        await dp_instance.feed_update(bot_instance, update)
    return {"ok": True}


@app.get("/health")
def health_check() -> dict[str, str]:
    return {"status": "ok"}


@app.get("/")
def serve_web_app() -> FileResponse:
    return FileResponse(static_dir / "index.html")


@app.get("/api/meta/statuses", response_model=list[StatusCatalogItem])
def get_statuses() -> list[StatusCatalogItem]:
    return status_catalog()


@app.get("/api/meta/marketplaces")
def get_marketplaces() -> list[dict[str, str]]:
    return marketplace_catalog()


@app.get("/api/orders", response_model=OrdersResponse)
def list_orders_endpoint(
    marketplace: Optional[Marketplace] = Query(default=None),
    search: Optional[str] = Query(default=None, max_length=128),
    limit: int = Query(default=200, ge=1, le=500),
    offset: int = Query(default=0, ge=0),
    session: Session = Depends(get_session),
) -> OrdersResponse:
    items, total = list_orders(
        session=session,
        marketplace=marketplace,
        search=search,
        limit=limit,
        offset=offset,
    )
    return OrdersResponse(items=items, total=total)


@app.get("/api/dashboard/{marketplace}", response_model=DashboardSummary)
def dashboard_marketplace_endpoint(
    marketplace: Marketplace,
    session: Session = Depends(get_session),
) -> DashboardSummary:
    return build_summary(session, marketplace)


@app.get("/api/dashboard")
def dashboard_all_endpoint(
    session: Session = Depends(get_session),
) -> list[DashboardSummary]:
    return [build_summary(session, marketplace) for marketplace in Marketplace]


@app.get("/api/settings", response_model=SettingsRead)
def settings_read_endpoint(
    session: Session = Depends(get_session),
) -> SettingsRead:
    return get_settings(session)


@app.put("/api/settings", response_model=SettingsRead)
def settings_update_endpoint(
    payload: SettingsUpdate,
    session: Session = Depends(get_session),
) -> SettingsRead:
    return save_settings(session, payload)


@app.post("/api/sync/run", response_model=SyncReport)
async def run_sync_endpoint() -> SyncReport:
    return await sync_orders_from_marketplaces()


@app.post("/api/admin/cleanup-srid")
def cleanup_srid_endpoint(session: Session = Depends(get_session)) -> dict:
    """Удаляет WB заказы с srid-номерами (содержащими буквы или точки)."""
    rows = session.execute(
        text("SELECT id, external_order_id FROM orders WHERE marketplace='wb'")
    ).fetchall()

    srid_ids = [
        row[0] for row in rows
        if any(c.isalpha() for c in str(row[1])) or '.' in str(row[1])
    ]

    if not srid_ids:
        return {"deleted_orders": 0, "deleted_events": 0, "message": "Нечего удалять"}

    placeholders = ",".join(str(i) for i in srid_ids)
    r1 = session.execute(text(f"DELETE FROM order_events WHERE order_id IN ({placeholders})"))
    r2 = session.execute(text(f"DELETE FROM orders WHERE id IN ({placeholders})"))
    session.commit()

    logger.info("cleanup-srid: удалено заказов=%s событий=%s", r2.rowcount, r1.rowcount)
    return {
        "deleted_orders": r2.rowcount,
        "deleted_events": r1.rowcount,
        "message": "Готово",
    }



@app.post("/api/admin/reset-wb")
def reset_wb_endpoint(session: Session = Depends(get_session)) -> dict:
    """Полная очистка всех WB заказов — для пересинхронизации с нуля."""
    r1 = session.execute(text("DELETE FROM order_events WHERE order_id IN (SELECT id FROM orders WHERE marketplace='wb')"))
    r2 = session.execute(text("DELETE FROM orders WHERE marketplace='wb'"))
    session.commit()
    logger.info("reset-wb: удалено заказов=%s событий=%s", r2.rowcount, r1.rowcount)
    return {"deleted_orders": r2.rowcount, "deleted_events": r1.rowcount, "message": "WB заказы удалены"}

@app.get("/api/export/orders.csv")
def export_orders_csv_endpoint(
    session: Session = Depends(get_session),
) -> StreamingResponse:
    rows = export_rows(session)
    columns = [
        "Маркетплейс",
        "Номер сборочного задания",
        "Текущий статус",
        "Дата текущего статуса",
        "История статусов",
    ]

    csv_buffer = io.StringIO()
    writer = csv.DictWriter(csv_buffer, fieldnames=columns, delimiter=";")
    writer.writeheader()
    writer.writerows(rows)

    filename = f"orders_export_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
    payload = io.BytesIO(csv_buffer.getvalue().encode("utf-8-sig"))
    return StreamingResponse(
        payload,
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@app.get("/api/export/orders.xlsx")
def export_orders_xlsx_endpoint(
    session: Session = Depends(get_session),
) -> StreamingResponse:
    rows = export_rows(session)
    columns = [
        "Маркетплейс",
        "Номер сборочного задания",
        "Текущий статус",
        "Дата текущего статуса",
        "История статусов",
    ]

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Заказы"
    sheet.append(columns)
    for row in rows:
        sheet.append([row[column] for column in columns])

    sheet.freeze_panes = "A2"
    widths = [18, 28, 30, 24, 90]
    for idx, width in enumerate(widths, start=1):
        column_letter = chr(64 + idx)
        sheet.column_dimensions[column_letter].width = width

    stream = io.BytesIO()
    workbook.save(stream)
    stream.seek(0)

    filename = f"orders_export_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
